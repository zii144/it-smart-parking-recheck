"""Excel import helpers for admin bulk data (locations, inspectors).

Templates use a header row plus an optional example/notes row; data rows follow.
Column headers accept zh-TW labels or English field names.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from .models import Inspector, Location
from .security import hash_password

IMPORT_TYPES = frozenset({"locations", "inspectors"})

LOCATION_HEADER_MAP = {
    "行政區": "district",
    "district": "district",
    "路段": "road",
    "road": "road",
    "停車格編號": "spot_no",
    "停車格": "spot_no",
    "spot_no": "spot_no",
}

INSPECTOR_HEADER_MAP = {
    "帳號": "username",
    "username": "username",
    "密碼": "password",
    "password": "password",
    "姓名": "display_name",
    "display_name": "display_name",
    "啟用權限": "has_permission",
    "has_permission": "has_permission",
}

LOCATION_REQUIRED = frozenset({"district", "road", "spot_no"})
INSPECTOR_REQUIRED = frozenset({"username", "password", "display_name"})


@dataclass
class ImportResult:
    import_type: str
    total_rows: int
    created: int
    skipped: int
    errors: list[dict[str, Any]] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "import_type": self.import_type,
            "total_rows": self.total_rows,
            "created": self.created,
            "skipped": self.skipped,
            "errors": self.errors,
        }


def _normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "是" if value else "否"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _is_example_row(cells: tuple[Any, ...]) -> bool:
    first = _cell_text(cells[0]) if cells else ""
    return first.startswith("範例") or first.startswith("示例") or first.lower().startswith("example")


def _row_is_empty(cells: tuple[Any, ...]) -> bool:
    return all(not _cell_text(c) for c in cells)


def _parse_permission(value: Any) -> tuple[bool | None, str | None]:
    text = _cell_text(value)
    if not text:
        return True, None
    lowered = text.lower()
    if lowered in {"是", "1", "true", "yes", "y", "啟用", "有"}:
        return True, None
    if lowered in {"否", "0", "false", "no", "n", "停用", "無"}:
        return False, None
    return None, f"啟用權限格式無效：{text}（請填 是/否）"


def parse_workbook_rows(content: bytes, import_type: str) -> tuple[list[tuple[int, dict[str, str]]], str | None]:
    """Return (rows with 1-based Excel row numbers, fatal parse error)."""
    header_map = LOCATION_HEADER_MAP if import_type == "locations" else INSPECTOR_HEADER_MAP
    required = LOCATION_REQUIRED if import_type == "locations" else INSPECTOR_REQUIRED

    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception:
        return [], "無法讀取 Excel 檔案，請確認格式為 .xlsx"

    ws = wb.active
    raw_rows = list(ws.iter_rows(values_only=True))
    if not raw_rows:
        return [], "檔案沒有資料列"

    header_row_idx = None
    col_keys: list[str | None] = []
    for idx, row in enumerate(raw_rows):
        if _row_is_empty(row):
            continue
        if _is_example_row(row):
            continue
        labels = [_normalize_header(c) for c in row]
        mapped = [header_map.get(label) for label in labels]
        if any(k in required for k in mapped if k):
            header_row_idx = idx
            col_keys = mapped
            break

    if header_row_idx is None:
        return [], "找不到有效的欄位標題列（請使用系統提供的範本）"

    missing = sorted(required - {k for k in col_keys if k})
    if missing:
        labels = {
            "district": "行政區",
            "road": "路段",
            "spot_no": "停車格編號",
            "username": "帳號",
            "password": "密碼",
            "display_name": "姓名",
        }
        return [], f"缺少必要欄位：{', '.join(labels[k] for k in missing)}"

    parsed: list[tuple[int, dict[str, str]]] = []
    for offset, row in enumerate(raw_rows[header_row_idx + 1 :], start=header_row_idx + 2):
        if _row_is_empty(row):
            continue
        if _is_example_row(row):
            continue
        record: dict[str, str] = {}
        for key, cell in zip(col_keys, row):
            if key:
                record[key] = _cell_text(cell)
        if any(record.get(k) for k in required):
            parsed.append((offset, record))

    if not parsed:
        return [], "沒有可匯入的資料列"

    return parsed, None


def build_template_workbook(import_type: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    if import_type == "locations":
        ws.title = "停車格匯入"
        ws.append(["行政區", "路段", "停車格編號"])
        ws.append(["範例：信義區", "松高路", "Z-001"])
    else:
        ws.title = "稽查員匯入"
        ws.append(["帳號", "密碼", "姓名", "啟用權限"])
        ws.append(["範例：insp99", "pass123", "測試員", "是"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def import_locations(db: Session, rows: list[tuple[int, dict[str, str]]]) -> ImportResult:
    result = ImportResult(import_type="locations", total_rows=len(rows), created=0, skipped=0)
    for row_num, row in rows:
        district = row.get("district", "").strip()
        road = row.get("road", "").strip()
        spot_no = row.get("spot_no", "").strip()
        if not district or not road or not spot_no:
            result.errors.append({"row": row_num, "message": "行政區、路段、停車格編號皆為必填"})
            continue

        existing = db.scalar(
            select(Location).where(
                Location.district == district,
                Location.road == road,
                Location.spot_no == spot_no,
            )
        )
        if existing:
            result.skipped += 1
            continue

        db.add(Location(district=district, road=road, spot_no=spot_no))
        try:
            db.commit()
            result.created += 1
        except IntegrityError:
            db.rollback()
            result.skipped += 1
    return result


def import_inspectors(db: Session, rows: list[tuple[int, dict[str, str]]]) -> ImportResult:
    result = ImportResult(import_type="inspectors", total_rows=len(rows), created=0, skipped=0)
    for row_num, row in rows:
        username = row.get("username", "").strip()
        password = row.get("password", "").strip()
        display_name = row.get("display_name", "").strip()
        if not username or not password or not display_name:
            result.errors.append({"row": row_num, "message": "帳號、密碼、姓名皆為必填"})
            continue

        has_permission, perm_err = _parse_permission(row.get("has_permission"))
        if perm_err:
            result.errors.append({"row": row_num, "message": perm_err})
            continue

        existing = db.scalar(select(Inspector).where(Inspector.username == username))
        if existing:
            result.skipped += 1
            continue

        db.add(
            Inspector(
                username=username,
                password=hash_password(password),
                display_name=display_name,
                has_permission=int(has_permission),
            )
        )
        try:
            db.commit()
            result.created += 1
        except IntegrityError:
            db.rollback()
            result.skipped += 1
    return result


def run_import(db: Session, import_type: str, content: bytes) -> tuple[ImportResult | None, str | None]:
    rows, parse_error = parse_workbook_rows(content, import_type)
    if parse_error:
        return None, parse_error
    if import_type == "locations":
        return import_locations(db, rows), None
    return import_inspectors(db, rows), None
